from datetime import datetime
import io
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import pdfplumber
import streamlit as st

# ==========================================
# 상수 및 설정
# ==========================================
MONTH_MAP = {
    "JAN": "01", "FEB": "02", "MAR": "03", "APR": "04",
    "MAY": "05", "JUN": "06", "JUL": "07", "AUG": "08",
    "SEP": "09", "OCT": "10", "NOV": "11", "DEC": "12",
}

# ==========================================
# 파싱 및 추출 로직
# ==========================================
def get_year_from_filename(filename):
    match = re.search(r"(20\d{2})", str(filename))
    if match:
        return match.group(1)
    return str(datetime.now().year)

def parse_date_text(date_text, filename):
    if not date_text:
        return None
    text = date_text.upper()
    year_from_path = get_year_from_filename(filename)
    
    match = re.search(r"(20\d{2})[-./](\d{1,2})[-./](\d{1,2})", text)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
        
    match = re.search(r"\b(\d{1,2})\s*([A-Z]{3})\s*(20\d{2})\b", text)
    if match:
        month_text = match.group(2)
        if month_text in MONTH_MAP:
            return f"{match.group(3)}-{MONTH_MAP[month_text]}-{int(match.group(1)):02d}"
            
    match = re.search(r"\b(\d{1,2})\s*([A-Z]{3})\b", text)
    if match:
        month_text = match.group(2)
        if month_text in MONTH_MAP:
            return f"{year_from_path}-{MONTH_MAP[month_text]}-{int(match.group(1)):02d}"
    return None

def extract_pdf_date(uploaded_file):
    try:
        with pdfplumber.open(uploaded_file) as pdf:
            if not pdf.pages:
                return None, ""
            page = pdf.pages[0]
            width, height = page.width, page.height
            crop_box = (width * 0.55, 0, width, height * 0.25)
            cropped = page.crop(crop_box)
            top_right_text = cropped.extract_text() or ""
            parsed_date = parse_date_text(top_right_text, uploaded_file.name)
            if not parsed_date:
                full_text = page.extract_text() or ""
                parsed_date = parse_date_text(full_text, uploaded_file.name)
            return parsed_date, top_right_text
    except Exception as e:
        return None, f"DATE 추출 오류: {e}"

def get_year_month(pdf_date):
    if not pdf_date:
        return "UNKNOWN_MONTH"
    try:
        dt = pd.to_datetime(pdf_date)
        return dt.strftime("%Y-%m")
    except Exception:
        return "UNKNOWN_MONTH"

def normalize_table(table):
    cleaned_rows = []
    for row in table:
        if not row:
            continue
        cleaned_row = ["" if cell is None else str(cell).strip() for cell in row]
        if any(cleaned_row):
            cleaned_rows.append(cleaned_row)
    if not cleaned_rows:
        return []
    max_cols = max(len(row) for row in cleaned_rows)
    return [row + [""] * (max_cols - len(row)) for row in cleaned_rows]

def extract_tables_from_pdf(uploaded_file, pdf_date, year_month):
    all_rows = []
    try:
        uploaded_file.seek(0)
        with pdfplumber.open(uploaded_file) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables()
                for table_no, table in enumerate(tables, start=1):
                    rows = normalize_table(table)
                    for row_index, row in enumerate(rows, start=1):
                        row_data = {
                            "PDF_DATE": pdf_date,
                            "YEAR_MONTH": year_month,
                            "source_folder": "Uploaded_File",
                            "source_file": uploaded_file.name,
                            "page": page_no,
                            "table_no": table_no,
                            "row_no": row_index,
                        }
                        for col_index, value in enumerate(row, start=1):
                            row_data[f"col_{col_index}"] = value
                        all_rows.append(row_data)
    except Exception as e:
        all_rows.append({
            "PDF_DATE": pdf_date,
            "YEAR_MONTH": year_month,
            "source_folder": "Uploaded_File",
            "source_file": uploaded_file.name,
            "page": "",
            "table_no": "",
            "row_no": "",
            "error": str(e),
        })
    return all_rows

def extract_visual_lines_from_page(page, y_tolerance=3):
    words = page.extract_words(x_tolerance=2, y_tolerance=3, use_text_flow=False, keep_blank_chars=False)
    if not words:
        return []
    words = sorted(words, key=lambda w: (round(float(w["top"]), 1), float(w["x0"])))
    lines, current_words, current_top = [], [], None
    for word in words:
        word_top = float(word["top"])
        if current_top is None:
            current_top, current_words = word_top, [word]
            continue
        if abs(word_top - current_top) <= y_tolerance:
            current_words.append(word)
        else:
            current_words = sorted(current_words, key=lambda w: float(w["x0"]))
            line_text = " ".join(w["text"] for w in current_words).strip()
            if line_text:
                lines.append({"visual_top": current_top, "text": line_text})
            current_top, current_words = word_top, [word]
    if current_words:
        current_words = sorted(current_words, key=lambda w: float(w["x0"]))
        line_text = " ".join(w["text"] for w in current_words).strip()
        if line_text:
            lines.append({"visual_top": current_top, "text": line_text})
    return lines

def extract_text_lines_from_pdf(uploaded_file, pdf_date, year_month):
    text_rows = []
    try:
        uploaded_file.seek(0)
        with pdfplumber.open(uploaded_file) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                visual_lines = extract_visual_lines_from_page(page)
                for line_no, item in enumerate(visual_lines, start=1):
                    line = item["text"].strip()
                    if line:
                        text_rows.append({
                            "PDF_DATE": pdf_date,
                            "YEAR_MONTH": year_month,
                            "source_folder": "Uploaded_File",
                            "source_file": uploaded_file.name,
                            "page": page_no,
                            "line_no": line_no,
                            "visual_top": item["visual_top"],
                            "text": line,
                        })
    except Exception as e:
        text_rows.append({
            "PDF_DATE": pdf_date,
            "YEAR_MONTH": year_month,
            "source_folder": "Uploaded_File",
            "source_file": uploaded_file.name,
            "page": "",
            "line_no": "",
            "visual_top": "",
            "text": "",
            "error": str(e),
        })
    return text_rows

def sort_dataframe_by_pdf_date(df):
    if df.empty or "PDF_DATE" not in df.columns:
        return df
    df["_sort_date"] = pd.to_datetime(df["PDF_DATE"], errors="coerce")
    sort_cols = ["_sort_date"] + [c for c in ["source_file", "page", "row_no", "line_no"] if c in df.columns]
    df = df.sort_values(by=sort_cols, ascending=True, na_position="last").drop(columns=["_sort_date"])
    return df

def clean_text(value):
    return "" if pd.isna(value) else str(value).replace("\n", " ").strip()

def extract_crew_names_from_line(line):
    text = clean_text(line).upper()
    role_name_pattern = re.compile(
        r"\b(?:RCAPT|CAPT|CPT|RF/O|RFO|F/O|FO|FR/P|FRP|S/O|SO|CKAIR|CK\s*AIR|DHD)\s*:?\s*([A-Z]\.\s*[A-Z][A-Z'’\-]+)",
        re.IGNORECASE,
    )
    names = []
    for match in role_name_pattern.finditer(text):
        name = re.sub(r"\s+", " ", match.group(1)).strip()
        name = re.sub(r"\s+(?:MR|MRS|MS)$", "", name)
        if name:
            names.append(name.upper())
    return names

def detect_new_flight_from_line(line):
    text = clean_text(line).upper()
    compact = re.sub(r"[^A-Z0-9]", "", text)
    if "PAX" in compact:
        match = re.search(r"\b([A-Z]{2,3}\d{2,5}[A-Z]?)\b", text)
        if match and not match.group(1).startswith("PAX"):
            return match.group(1)
    match = re.match(r"^\s*FX\s+(\d{3,5})-\d{1,3}\b", text)
    if match:
        return f"FX{match.group(1)}"
    match = re.match(r"^\s*(\d{3,5})-\d{1,3}\b", text)
    if match:
        return f"FX{match.group(1)}"
    return None

def extract_route_raw_from_line(line):
    text = re.sub(r"\s+", " ", clean_text(line).upper())
    place_pattern = r"(?:GRAND\s+HYATT|ICN\s+HYATT|SHERATON|T1|T2|SPOT|RAMP|여객)"
    match = re.search(rf"({place_pattern})\s*-\s*({place_pattern})", text, re.IGNORECASE)
    if match:
        return f"{re.sub(r' ', ' ', match.group(1).strip())}-{re.sub(r' ', ' ', match.group(2).strip())}"
    return ""

def is_noise_line(line):
    text = clean_text(line).upper()
    noise_starts = [
        "P :", "S :", "T :", "NOTE", "DAILY WORK PLAN", "CIQ ACTIVITY",
        "REMARKS", "FLT NBR", "CREW LIST", "ALERT", "PICK-UP", "RESER",
        "RESERVATION", "G/T", "D3 DATE", "DATE :",
    ]
    return any(text.startswith(k) for k in noise_starts)

def deduplicate_keep_order(values):
    seen, result = set(), []
    for v in values:
        k = v.strip().upper()
        if k and k not in seen:
            seen.add(k)
            result.append(v.strip())
    return result

def sort_text_lines_by_visual_position(df):
    if df.empty:
        return df
    df = df.copy()
    df["_sort_date"] = pd.to_datetime(df["PDF_DATE"], errors="coerce")
    for col in ["visual_top", "page", "line_no"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    sort_cols = ["_sort_date", "source_file", "page", "visual_top", "line_no"]
    return df.sort_values(by=sort_cols, ascending=True, na_position="last", kind="mergesort").drop(columns=["_sort_date"])

def sort_processed_crew_dataframe(df):
    if df.empty or "PDF_DATE" not in df.columns:
        return df
    df = df.copy()
    df["_sort_date"] = pd.to_datetime(df["PDF_DATE"], errors="coerce")
    for col in ["PAGE", "START_TOP", "START_LINE", "PDF_ORDER"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    sort_cols = ["_sort_date"] + [c for c in ["SOURCE_FILE", "PAGE", "START_TOP", "START_LINE", "PDF_ORDER"] if c in df.columns]
    return df.sort_values(by=sort_cols, ascending=True, na_position="last", kind="mergesort").drop(columns=["_sort_date"])

def extract_flight_time(lines):
    time_pickup, time_eta_etd = "", ""
    for line in lines:
        text = str(line).upper()
        match_eta = re.search(r"(ETA|ETD)\s*(\d{2}:\d{2})", text)
        if match_eta:
            time_eta_etd = match_eta.group(2).strip()

        all_times = re.findall(r"\d{2}:\d{2}", text)
        if all_times:
            if match_eta:
                eta_time_str = match_eta.group(2).strip()
                other_times = [t for t in all_times if t.strip() != eta_time_str]
                if other_times:
                    time_pickup = other_times[-1].strip()
            else:
                time_pickup = all_times[-1].strip()

    return time_pickup if time_pickup else time_eta_etd

def build_processed_crew_sheet_from_text(df_text):
    if df_text.empty:
        return pd.DataFrame()
    for col in ["PDF_DATE", "YEAR_MONTH", "source_file", "source_folder", "page", "line_no", "visual_top", "text"]:
        if col not in df_text.columns:
            df_text[col] = ""
    df = sort_text_lines_by_visual_position(df_text.copy())
    flight_blocks, current_block, last_doc_key, group_seq_by_doc = [], None, None, {}

    for _, row in df.iterrows():
        pdf_date, source_file, line = row.get("PDF_DATE", ""), row.get("source_file", ""), clean_text(row.get("text", ""))
        if not line:
            continue
        doc_key = (str(pdf_date), str(source_file))
        if doc_key != last_doc_key:
            current_block, last_doc_key = None, doc_key
        if doc_key not in group_seq_by_doc:
            group_seq_by_doc[doc_key] = 0
        if is_noise_line(line):
            continue

        flight_no = detect_new_flight_from_line(line)
        names = extract_crew_names_from_line(line)
        route = extract_route_raw_from_line(line)

        if flight_no:
            group_seq_by_doc[doc_key] += 1
            current_block = {
                "PDF_DATE": pdf_date,
                "YEAR_MONTH": row.get("YEAR_MONTH", ""),
                "PDF_ORDER": group_seq_by_doc[doc_key],
                "PAGE": row.get("page", ""),
                "START_LINE": row.get("line_no", ""),
                "START_TOP": row.get("visual_top", ""),
                "FLIGHT_NO": flight_no,
                "SOURCE_FILE": source_file,
                "SOURCE_FOLDER": row.get("source_folder", ""),
                "lines_info": [],
            }
            flight_blocks.append(current_block)
        if current_block:
            current_block["lines_info"].append({"text": line, "crews": names, "route": route})

    result_rows = []
    for block in flight_blocks:
        lines_info = block["lines_info"]
        for i, info in enumerate(lines_info):
            text_upper = info["text"].replace(" ", "").upper()
            if "당일출국" in text_upper or "SELF" in text_upper:
                if info["crews"]:
                    info["crews"] = []
                else:
                    for j in range(i - 1, -1, -1):
                        if lines_info[j]["crews"]:
                            lines_info[j]["crews"] = []
                            break

        sub_groups = []
        for info in lines_info:
            c, r = info["crews"], info["route"]
            if not sub_groups:
                sub_groups.append({"route": r, "crews": list(c), "raw_lines": [info["text"]]})
                continue
            if r:
                if not sub_groups[-1]["route"] or sub_groups[-1]["route"] == r:
                    sub_groups[-1]["route"] = r
                    sub_groups[-1]["crews"].extend(c)
                    sub_groups[-1]["raw_lines"].append(info["text"])
                else:
                    sub_groups.append({"route": r, "crews": list(c), "raw_lines": [info["text"]]})
            else:
                sub_groups[-1]["crews"].extend(c)
                sub_groups[-1]["raw_lines"].append(info["text"])

        for sub in sub_groups:
            unique_names = deduplicate_keep_order(sub["crews"])
            if not unique_names:
                continue
            ex_time = extract_flight_time(sub["raw_lines"]) or extract_flight_time([info["text"] for info in lines_info])
            result_rows.append({
                "PDF_DATE": block["PDF_DATE"],
                "FLIGHT_NO": block["FLIGHT_NO"],
                "CREW_NAMES": " / ".join(unique_names),
                "ROUTE_RAW": sub["route"],
                "CREW_COUNT": len(unique_names),
                "TIME": ex_time,
                "YEAR_MONTH": block["YEAR_MONTH"],
                "PDF_ORDER": block["PDF_ORDER"],
                "PAGE": block["PAGE"],
                "START_LINE": block["START_LINE"],
                "START_TOP": block.get("START_TOP", ""),
                "RAW_FLIGHT_LINE": " \n ".join(sub["raw_lines"]),
                "SOURCE_FILE": block["SOURCE_FILE"],
                "SOURCE_FOLDER": block["SOURCE_FOLDER"],
            })

    result_df = pd.DataFrame(result_rows)
    if result_df.empty:
        return result_df
    pref_cols = [
        "PDF_DATE", "FLIGHT_NO", "CREW_NAMES", "ROUTE_RAW", "CREW_COUNT",
        "TIME", "YEAR_MONTH", "PDF_ORDER", "PAGE", "START_LINE", "START_TOP",
        "RAW_FLIGHT_LINE", "SOURCE_FILE", "SOURCE_FOLDER",
    ]
    result_df = result_df[[c for c in pref_cols if c in result_df.columns] + [c for c in result_df.columns if c not in pref_cols]]
    return sort_processed_crew_dataframe(result_df)

def create_time_schedule_sheet(df):
    if df.empty or "TIME" not in df.columns:
        return pd.DataFrame()
    res_df = df.copy()

    res_df["_sort_time"] = pd.to_datetime(res_df["TIME"], format="%H:%M", errors="coerce")
    res_df = res_df.sort_values(by=["PDF_DATE", "_sort_time", "FLIGHT_NO"], na_position="last")

    target_cols = ["FLIGHT_NO", "CREW_NAMES", "ROUTE_RAW", "CREW_COUNT", "TIME"]
    final_rows, prev_hour, prev_date = [], None, None
    idx = 1

    for _, row in res_df.iterrows():
        curr_time, curr_date = str(row.get("TIME", "")).strip(), row.get("PDF_DATE", "")
        curr_hour = ""
        if curr_time:
            match = re.search(r"(\d{1,2}):\d{2}", curr_time)
            if match:
                curr_hour = match.group(1)

        if prev_hour is not None and curr_time != "":
            if curr_hour != prev_hour or curr_date != prev_date:
                blank_row = {"INDEX": ""}
                for col in target_cols:
                    blank_row[col] = ""
                final_rows.append(blank_row)

        data_row = {"INDEX": idx}
        for col in target_cols:
            data_row[col] = row.get(col, "")

        final_rows.append(data_row)
        idx += 1
        prev_hour, prev_date = curr_hour, curr_date

    return pd.DataFrame(final_rows)

# ==========================================
# 엑셀 스타일링 및 생성 (메모리 버퍼 대응)
# ==========================================
def style_excel_buffer(buffer):
    wb = load_workbook(buffer)
    header_fill, header_font = PatternFill("solid", fgColor="E2E9F3"), Font(color="333333", bold=True)
    idx_fill, idx_font = PatternFill("solid", fgColor="F2F4F8"), Font(color="333333", bold=True)

    thin_border = Side(style="thin", color="000000")
    border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    no_border = Border()
    no_fill = PatternFill(fill_type=None)

    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue
        ws.freeze_panes = "A2"

        gap_cols = []
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=col_idx).value
            if val in [" ", "  "]:
                gap_cols.append(col_idx)

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            if col_idx in gap_cols:
                cell.value = None
                cell.border = no_border
                cell.fill = no_fill
            else:
                cell.fill, cell.font, cell.alignment, cell.border = (
                    header_fill,
                    header_font,
                    Alignment(horizontal="center", vertical="center"),
                    border,
                )

        pdf_date_col, index_col = None, None
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=col_idx).value
            if val == "PDF_DATE":
                pdf_date_col = col_idx
            if val == "INDEX":
                index_col = col_idx

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                if cell.column in gap_cols:
                    cell.border = no_border
                    cell.fill = no_fill
                else:
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        target_col = index_col if index_col else pdf_date_col
        if target_col:
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=target_col)
                if str(cell.value).strip() and cell.column not in gap_cols:
                    cell.fill, cell.font = idx_fill, idx_font

        if gap_cols:
            last_filter_col = get_column_letter(gap_cols[0] - 1)
            ws.auto_filter.ref = f"A1:{last_filter_col}{ws.max_row}"
        else:
            ws.auto_filter.ref = ws.dimensions

        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            if col_idx in gap_cols:
                ws.column_dimensions[col_letter].width = 4
                continue

            max_length = 0
            for row_idx in range(1, min(ws.max_row, 300) + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val:
                    val_str = str(val)
                    lines = val_str.split("\n")
                    longest_line = max(lines, key=len)
                    max_length = max(max_length, len(longest_line) * 1.3)

            ws.column_dimensions[col_letter].width = min(max(max_length + 2, 10), 100)

        ws.row_dimensions[1].height = 22

    output_buffer = io.BytesIO()
    wb.save(output_buffer)
    output_buffer.seek(0)
    return output_buffer

def generate_excel_file(df_tables, df_text, df_log):
    if df_log.empty:
        return None

    df_processed_crew = build_processed_crew_sheet_from_text(df_text)
    df_time_schedule = create_time_schedule_sheet(df_processed_crew)

    if not df_processed_crew.empty and "PDF_DATE" in df_processed_crew.columns:
        df_processed_crew.insert(0, "INDEX", range(1, len(df_processed_crew) + 1))
        df_processed_crew = df_processed_crew.drop(columns=["PDF_DATE"])

    if not df_processed_crew.empty and "TIME" in df_processed_crew.columns:
        time_idx = df_processed_crew.columns.get_loc("TIME")
        df_processed_crew.insert(time_idx + 1, " ", "")
        df_processed_crew.insert(time_idx + 2, "  ", "")

    df_tables = sort_dataframe_by_pdf_date(df_tables)
    df_text = sort_dataframe_by_pdf_date(df_text)
    df_log = sort_dataframe_by_pdf_date(df_log)

    raw_buffer = io.BytesIO()
    with pd.ExcelWriter(raw_buffer, engine="openpyxl") as writer:
        if not df_time_schedule.empty:
            df_time_schedule.to_excel(writer, sheet_name="시간대별", index=False)
        else:
            pd.DataFrame([{"Message": "No data"}]).to_excel(writer, sheet_name="시간대별", index=False)

        if not df_processed_crew.empty:
            df_processed_crew.to_excel(writer, sheet_name="PDF_순서_SKD", index=False)
        else:
            pd.DataFrame([{"Message": "No data"}]).to_excel(writer, sheet_name="PDF_순서_SKD", index=False)

        if not df_tables.empty:
            df_tables.to_excel(writer, sheet_name="TABLE_DATA_SORTED", index=False)
        else:
            pd.DataFrame([{"Message": "No data"}]).to_excel(writer, sheet_name="TABLE_DATA_SORTED", index=False)

        if not df_text.empty:
            df_text.to_excel(writer, sheet_name="TEXT_LINES_SORTED", index=False)
        else:
            pd.DataFrame([{"Message": "No data"}]).to_excel(writer, sheet_name="TEXT_LINES_SORTED", index=False)

        df_log.to_excel(writer, sheet_name="LOG", index=False)

    raw_buffer.seek(0)
    return style_excel_buffer(raw_buffer), df_time_schedule, df_processed_crew

def extract_date_from_filename(filename):
    match = re.search(r"(\d{4}-\d{2}-\d{2})", filename)
    return match.group(1) if match else None

# ==========================================
# Streamlit 웹 UI 부분
# ==========================================
from pathlib import Path
import pandas as pd
import streamlit as st


def main():
    st.set_page_config(
        page_title="CIQ Activity PDF 파서", page_icon="✈️", layout="wide"
    )

    st.title("✈️ Airman Daily Activity 엑셀 변환기 ✈️")
    st.markdown("PDF 파일을 업로드하면 **시간대별로 엑셀 문서로 생성**합니다.")
    st.divider()

    # Subheader 가운데 정렬 스타일 적용
    st.markdown(
        """
        <style>
        /* Subheader 중앙 정렬 */
        [data-testid="stSubheader"] {
            text-align: center;
            width: 100%;
        }
        
        /* 파일 업로더 안내 문구 글씨 크기 및 두께 키우기 */
        [data-testid="stFileUploader"] label p {
            font-size: 18px !important;      /* 기본보다 크게 (원하는 크기로 숫자 변경 가능) */
            font-weight: bold !important;    /* 글씨 두껍게 */
            color: #1F2937 !important;       /* 글씨 색상 (진한 회색) */
        }
        
        /* 일반 버튼(st.button)과 다운로드 버튼(st.download_button) 공통 로고 색상 스타일 */
        div.stButton > button, div.stDownloadButton > button {
            height: 52px !important;
            font-size: 17px !important;
            font-weight: bold !important;
            border-radius: 8px !important;
            width: 100% !important;
            background-color: #0038A8 !important;  /* 로고 파란색 (#0038A8) */
            color: #FFFFFF !important;              /* 글자색 흰색 */
            border: none !important;
            transition: background-color 0.2s ease, transform 0.1s ease;
        }

        /* 마우스를 올렸을 때 (Hover) */
        div.stButton > button:hover, div.stDownloadButton > button:hover {
            background-color: #002878 !important;
            color: #FFFFFF !important;
        }

        /* 버튼을 누를 때 (Active / Focus) */
        div.stButton > button:active, div.stDownloadButton > button:active,
        div.stButton > button:focus, div.stDownloadButton > button:focus {
            background-color: #001F58 !important;
            color: #FFFFFF !important;
            box-shadow: none !important;
        }
        </style>
    """,
        unsafe_allow_html=True,
    )

    uploaded_file = st.file_uploader(
        "분석할 PDF 파일을 선택하거나 드래그해 주세요.", type=["pdf"]
    )

    if uploaded_file is not None:
        st.success(f"파일 선택 완료: `{uploaded_file.name}`")

        # 새로 업로드된 파일이 바뀌면 이전 분석 데이터 초기화
        if (
            "last_uploaded_file" not in st.session_state
            or st.session_state["last_uploaded_file"] != uploaded_file.name
        ):
            st.session_state["last_uploaded_file"] = uploaded_file.name
            st.session_state["excel_data"] = None

        # [버튼 가로 배치] 시작 버튼과 다운로드 버튼을 나란히 놓기 위한 2개 컬럼 생성
        col1, spacer, col2 = st.columns([2, 0.2, 2], gap="large")

        with col1:
            if st.button(
                "🚀 분석 및 엑셀 생성 시작",
                type="primary",
                use_container_width=True,
            ):
                with st.spinner("PDF 데이터를 파싱하는 중입니다..."):
                    pdf_date, date_raw_text = extract_pdf_date(uploaded_file)
                    if not pdf_date:
                        pdf_date = extract_date_from_filename(
                            uploaded_file.name
                        )

                    # 💡 [수정 위치] pdf_date를 구한 후 항상 year_month를 생성하도록 if문 밖으로 이동
                    year_month = get_year_month(pdf_date)

                    extracted_tables = extract_tables_from_pdf(
                        uploaded_file, pdf_date, year_month
                    )
                    extracted_texts = extract_text_lines_from_pdf(
                        uploaded_file, pdf_date, year_month
                    )

                    df_tables = pd.DataFrame(extracted_tables)
                    df_text = pd.DataFrame(extracted_texts)
                    df_log = pd.DataFrame([
                        {
                            "PDF_DATE": pdf_date,
                            "YEAR_MONTH": year_month,
                            "source_folder": "Uploaded_File",
                            "source_file": uploaded_file.name,
                            "date_raw_text_top_right": date_raw_text,
                            "table_rows": len(extracted_tables),
                            "text_lines": len(extracted_texts),
                        }
                    ])

                    (
                        excel_data,
                        df_time_schedule,
                        df_processed_crew,
                    ) = generate_excel_file(df_tables, df_text, df_log)

                    # 결과물을 세션(Session State)에 저장
                    st.session_state["excel_data"] = excel_data
                    st.session_state["df_time_schedule"] = df_time_schedule
                    st.session_state["df_processed_crew"] = df_processed_crew
                    st.session_state["output_filename"] = (
                        f"{Path(uploaded_file.name).stem}_변환결과.xlsx"
                    )

                st.toast("파싱이 성공적으로 완료되었습니다!", icon="✅")

        with col2:
            # 분석 결과가 세션에 저장되어 있을 때만 바로 옆에 다운로드 버튼 표시
            if st.session_state.get("excel_data") is not None:
                st.download_button(
                    label="📥 엑셀 파일 다운로드",
                    data=st.session_state["excel_data"],
                    file_name=st.session_state["output_filename"],
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                    use_container_width=True,
                )

        # 분석 완료 후 세션에 데이터가 존재하면 하단에 넓은 표 출력
        if st.session_state.get("excel_data") is not None:
            st.markdown("---")
            st.subheader("📊 결과 리스트보기")
            tab1, tab2 = st.tabs(["시간대별 스케줄", "PDF 순서 SKD"])
            with tab1:
                st.dataframe(
                    st.session_state["df_time_schedule"],
                    use_container_width=True,
                    height=600,
                )
            with tab2:
                st.dataframe(
                    st.session_state["df_processed_crew"],
                    use_container_width=True,
                    height=600,
                )


if __name__ == "__main__":
    main()